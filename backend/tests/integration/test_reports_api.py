"""The reports, as the Console reads them and as the day writes them.

Behavioural: trades go in through the store the engine uses, and the
claims are about the two reports' rows, their totals, the files they
export, and the end-of-day rows the REPORTS phase leaves behind.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from collections.abc import AsyncIterator
from typing import Any

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from garuda.api import envelope
from garuda.api.routers import reports
from garuda.composition.reports import write_daily_pnl
from garuda.composition.venues import venues_from
from garuda.core.clock import ReplayClock
from garuda.domain import Currency, Money, ProductType
from garuda.domain.client import TradingClientId
from garuda.domain.enums import Direction
from garuda.domain.instrument import InstrumentId
from garuda.domain.trade import Trade, TradeId
from garuda.domain.trade_state import TradeExitReason
from garuda.persistence import models
from garuda.persistence.uow import UnitOfWork

pytestmark = pytest.mark.integration

#: 09:30 in Kolkata on Monday the 12th.
NOW = dt.datetime(2026, 10, 12, 4, 0, tzinfo=dt.UTC)
MONDAY_MORNING = dt.datetime(2026, 10, 12, 4, 30, tzinfo=dt.UTC)
#: 01:30 IST on Tuesday the 13th.
TUESDAY_SMALL_HOURS = dt.datetime(2026, 10, 12, 20, 0, tzinfo=dt.UTC)
APPA = TradingClientId("appa")
AMMA = TradingClientId("amma")
CALL = InstrumentId("NSE:NIFTY26OCT25000CE")


def rupees(value: str) -> Money:
    return Money.of(value, Currency.INR)


def a_trade(
    trade_id: str,
    *,
    client: TradingClientId = APPA,
    strategy: str = "straddle",
    product: ProductType = ProductType.MIS,
    is_paper: bool = False,
) -> Trade:
    return Trade(
        id=TradeId(trade_id),
        trading_client=client,
        instrument=CALL,
        strategy=strategy,
        direction=Direction.SHORT,
        product=product,
        quantity=75,
        is_paper=is_paper,
    )


def closed(trade: Trade, entry: str, exit_price: str, *, at: dt.datetime = MONDAY_MORNING) -> Trade:
    return trade.with_entry_fill(trade.quantity, rupees(entry), at).closed(
        rupees(exit_price), TradeExitReason.TARGET, at
    )


async def seed(sessions: async_sessionmaker[AsyncSession]) -> None:
    async with sessions() as session:
        session.add(
            models.ExchangesRow(
                exchange_code="NSE",
                exchange_name="NSE",
                timezone="Asia/Kolkata",
                currency="INR",
                segments="EQUITY,FNO",
                market_open=dt.time(9, 15),
                market_close=dt.time(15, 30),
            )
        )
        for client_id, name in ((APPA.value, "Appa"), (AMMA.value, "Amma")):
            session.add(
                models.TradingClientRow(
                    id=client_id,
                    display_name=name,
                    broker="ZERODHA",
                    client_id=client_id.upper(),
                    enabled=True,
                    created_at=NOW,
                    updated_at=NOW,
                )
            )
        await session.commit()
    async with UnitOfWork(sessions) as uow:
        for trade in (
            closed(a_trade("a1"), "120", "100"),
            closed(a_trade("a2", strategy="strangle"), "100", "110"),
            closed(a_trade("a3", client=AMMA), "120", "110"),
            closed(a_trade("b1"), "120", "100", at=TUESDAY_SMALL_HOURS),
            closed(a_trade("p1", is_paper=True), "120", "100"),
            a_trade("open").with_entry_fill(75, rupees("1"), MONDAY_MORNING),
        ):
            await uow.repositories.live_trades.upsert_trade(trade, NOW)
        await uow.session.commit()


@pytest.fixture
async def client(
    session_factory: async_sessionmaker[AsyncSession],
) -> AsyncIterator[httpx.AsyncClient]:
    await seed(session_factory)
    app = FastAPI()
    envelope.install(app)
    app.include_router(reports.build(session_factory, ReplayClock(NOW)))
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as http:
        yield http


async def data(http: httpx.AsyncClient, path: str, **params: Any) -> dict[str, Any]:
    response = await http.get(path, params=params)
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()["data"]
    return body


class TestTheTradesReport:
    async def test_every_trade_in_the_range_with_its_totals(
        self, client: httpx.AsyncClient
    ) -> None:
        body = await data(client, "/api/reports/trades", fromDate="2026-10-12", toDate="2026-10-13")
        assert len(body["rows"]) == 6
        [a1] = [r for r in body["rows"] if r["tradeId"] == "a1"]
        assert a1["tradingClient"] == "Appa"
        assert a1["tradingSymbol"] == "NIFTY26OCT25000CE"
        assert a1["grossPnl"] == "1500"
        assert a1["charges"] is None, "not computed, never a zero"
        assert a1["tradingDay"] == "2026-10-12"
        [b1] = [r for r in body["rows"] if r["tradeId"] == "b1"]
        assert b1["tradingDay"] == "2026-10-13", "the venue's day"
        assert body["totals"]["trades"] == 6
        assert body["totals"]["closed"] == 5
        assert body["totals"]["grossPnl"] == "4500"
        assert body["totals"]["winRate"] == "80.00"

    async def test_the_filters_narrow_it(self, client: httpx.AsyncClient) -> None:
        mine = await data(
            client, "/api/reports/trades", fromDate="2026-10-12", tradingClientId="amma"
        )
        assert [r["tradeId"] for r in mine["rows"]] == ["a3"]
        paper = await data(client, "/api/reports/trades", fromDate="2026-10-12", paper="paper")
        assert [r["tradeId"] for r in paper["rows"]] == ["p1"]
        still_open = await data(
            client, "/api/reports/trades", fromDate="2026-10-12", state="ACTIVE"
        )
        assert [r["tradeId"] for r in still_open["rows"]] == ["open"]
        monday_only = await data(
            client, "/api/reports/trades", fromDate="2026-10-12", toDate="2026-10-12"
        )
        assert "b1" not in [r["tradeId"] for r in monday_only["rows"]]
        assert (
            await client.get("/api/reports/trades", params={"fromDate": "soon"})
        ).status_code == 400
        assert (
            await client.get("/api/reports/trades", params={"product": "LEVERAGED"})
        ).status_code == 400


class TestTheDailyPnl:
    async def test_a_day_for_whatever_was_selected(self, client: httpx.AsyncClient) -> None:
        body = await data(
            client,
            "/api/reports/daily-pnl",
            fromDate="2026-10-12",
            toDate="2026-10-13",
            paper="live",
        )
        rows = {r["tradingDay"]: r for r in body["rows"]}
        assert rows["2026-10-12"]["trades"] == 3
        assert rows["2026-10-12"]["grossPnl"] == "1500"
        assert rows["2026-10-12"]["netPnl"] == "1500"
        assert rows["2026-10-12"]["strategy"] is None
        assert rows["2026-10-13"]["grossPnl"] == "1500"
        assert body["tradingDays"] == 2
        assert body["rows"][0]["tradingDay"] == "2026-10-13", "newest first"

    async def test_one_strategy_or_strategy_wise(self, client: httpx.AsyncClient) -> None:
        one = await data(
            client,
            "/api/reports/daily-pnl",
            fromDate="2026-10-12",
            toDate="2026-10-12",
            strategy="strangle",
            paper="live",
        )
        [row] = one["rows"]
        assert row["grossPnl"] == "-750"
        split = await data(
            client,
            "/api/reports/daily-pnl",
            fromDate="2026-10-12",
            toDate="2026-10-12",
            byStrategy="true",
            paper="live",
        )
        assert {r["strategy"]: r["grossPnl"] for r in split["rows"]} == {
            "straddle": "2250",
            "strangle": "-750",
        }
        by_account = await data(
            client,
            "/api/reports/daily-pnl",
            fromDate="2026-10-12",
            toDate="2026-10-12",
            byAccount="true",
            paper="live",
        )
        assert {r["tradingClient"]: r["grossPnl"] for r in by_account["rows"]} == {
            "Appa": "750",
            "Amma": "750",
        }


class TestTheFiles:
    async def test_csv_carries_the_same_rows(self, client: httpx.AsyncClient) -> None:
        response = await client.get(
            "/api/reports/daily-pnl/export",
            params={
                "fromDate": "2026-10-12",
                "toDate": "2026-10-13",
                "paper": "live",
                "format": "csv",
            },
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert (
            'filename="daily_pnl_2026-10-12_2026-10-13.csv"'
            in response.headers["content-disposition"]
        )
        rows = list(csv.reader(io.StringIO(response.text)))
        assert rows[0][:2] == ["Trading day", "Account"]
        assert rows[1][0] == "2026-10-13"
        assert rows[2][8] == "1500", "gross P&L on Monday"

    async def test_excel_opens_with_the_same_rows(self, client: httpx.AsyncClient) -> None:
        response = await client.get(
            "/api/reports/trades/export",
            params={"fromDate": "2026-10-12", "toDate": "2026-10-13", "format": "xlsx"},
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        assert sheet is not None
        header = [c.value for c in sheet[1]]
        assert header[0] == "Trade"
        assert sheet.max_row == 7, "a header and six trades"
        gross = {row[0].value: row[17].value for row in sheet.iter_rows(min_row=2)}
        assert gross["a1"] == 1500

    async def test_an_unknown_format_is_refused(self, client: httpx.AsyncClient) -> None:
        response = await client.get("/api/reports/trades/export", params={"format": "pdf"})
        assert response.status_code == 400


class TestTheDaysRecord:
    async def test_the_reports_phase_writes_the_day_and_rewrites_it_the_same(
        self, session_factory: async_sessionmaker[AsyncSession]
    ) -> None:
        await seed(session_factory)
        nse = venues_from(
            [
                models.ExchangesRow(
                    exchange_code="NSE",
                    exchange_name="NSE",
                    timezone="Asia/Kolkata",
                    currency="INR",
                    segments="EQUITY,FNO",
                    market_open=dt.time(9, 15),
                    market_close=dt.time(15, 30),
                    is_active=True,
                )
            ],
            [],
        ).exchanges["NSE"]

        written = await write_daily_pnl(session_factory, nse, dt.date(2026, 10, 12))
        again = await write_daily_pnl(session_factory, nse, dt.date(2026, 10, 12))

        assert written == 4, "appa straddle, appa strangle, amma straddle, appa paper"
        assert again == 4
        async with session_factory() as session:
            rows = list((await session.execute(select(models.EodPnlReportsRow))).scalars())
        assert len(rows) == 4, "upserted, not doubled"
        by_key = {(r.trading_client_id, r.strategy_name, r.is_paper_trading): r for r in rows}
        appa = by_key[("appa", "straddle", False)]
        assert appa.date == dt.date(2026, 10, 12)
        assert appa.product == "MIS"
        assert str(appa.profit_loss) == "1500.000000"
        assert appa.charges is None
        assert str(appa.net_profit_loss) == "1500.000000"
        assert ("appa", "strangle", False) in by_key
        assert by_key[("appa", "straddle", True)].is_paper_trading is True
