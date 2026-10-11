"""Turning a report into a file the browser saves: CSV, or Excel.

CSV costs nothing. Excel is the owner's ask and is written with openpyxl,
the one dependency the reports add. Both take the same rows -- a header
and a list of cell lists -- so a report decides its columns once.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from fastapi import HTTPException
from fastapi.responses import Response

CSV_TYPE = "text/csv; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _cell(value: Any) -> Any:
    """What a cell holds. Money stays exact as text in CSV; Excel gets a number
    it can sum, from the exact decimal rather than a float on the way."""
    if isinstance(value, Decimal):
        return value
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    return value


def as_csv(header: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    for row in rows:
        writer.writerow([_cell(v) for v in row])
    return buffer.getvalue().encode("utf-8")


def as_xlsx(header: Sequence[str], rows: Sequence[Sequence[Any]], *, sheet: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    page = book.active
    assert page is not None
    page.title = sheet[:31]
    page.append(list(header))
    for cell in page[1]:
        cell.font = Font(bold=True)
    for row in rows:
        page.append([_cell(v) for v in row])
    page.freeze_panes = "A2"
    for index, column in enumerate(page.columns, start=1):
        widest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        page.column_dimensions[get_column_letter(index)].width = min(max(10, widest + 2), 48)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def file_response(
    fmt: str, *, name: str, header: Sequence[str], rows: Sequence[Sequence[Any]]
) -> Response:
    """The file, as the format asks: ``csv`` or ``xlsx``."""
    wanted = (fmt or "csv").strip().lower()
    if wanted == "csv":
        body, media, suffix = as_csv(header, rows), CSV_TYPE, "csv"
    elif wanted in ("xlsx", "excel"):
        body, media, suffix = as_xlsx(header, rows, sheet=name), XLSX_TYPE, "xlsx"
    else:
        raise HTTPException(status_code=400, detail=f"{fmt!r} is not an export format; csv or xlsx")
    return Response(
        content=body,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{name}.{suffix}"'},
    )
